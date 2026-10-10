
from pathlib import Path

import pymupdf
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def extract_text(file_path: str | Path) -> str:
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError("The uploaded document was not found.")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF, TXT, and DOCX files are supported.")

    if path.stat().st_size > MAX_FILE_SIZE:
        raise ValueError("The document exceeds the 10 MB limit.")

    if extension == ".txt":
        text = path.read_text(encoding="utf-8-sig")

    elif extension == ".pdf":
        with pymupdf.open(path) as document:
            if document.is_encrypted:
                raise ValueError(
                    "Password-protected PDFs are not supported."
                )

            text = "\n".join(page.get_text() for page in document)

    else:
        document = Document(path)

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        table_text = [
            " | ".join(
                cell.text.strip() for cell in row.cells
            )
            for table in document.tables
            for row in table.rows
        ]

        text = "\n".join(paragraphs + table_text)

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text found in the document."
        )

    return text
